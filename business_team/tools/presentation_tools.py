"""
Presentation tools for the Business Analyst Agent.

Generates PowerPoint presentations with charts, tables, and formatted slides.
"""

import json
from datetime import datetime
from pathlib import Path
from typing import Any

from business_team import config

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.chart import XL_CHART_TYPE

    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False


class PresentationTools:
    """Tools for creating PowerPoint presentations."""

    def __init__(self):
        self.output_dir = config.OUTPUT_DIR / "presentations"
        self.output_dir.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def get_tool_definitions() -> list[dict]:
        return [
            {
                "name": "create_presentation",
                "description": "Create a new PowerPoint presentation with title slide and content slides.",
                "input_schema": {
                    "type": "object",
                    "properties": {
                        "title": {
                            "type": "string",
                            "description": "Presentation title",
                        },
                        "subtitle": {
                            "type": "string",
                            "description": "Presentation subtitle",
                        },
                        "slides": {
                            "type": "array",
                            "description": "Array of slide objects",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "title": {"type": "string"},
                                    "content": {
                                        "type": "string",
                                        "description": "Slide body text (use \\n for line breaks)",
                                    },
                                    "slide_type": {
                                        "type": "string",
                                        "enum": [
                                            "title",
                                            "content",
                                            "two_column",
                                            "table",
                                            "kpi",
                                        ],
                                    },
                                    "data": {
                                        "type": "object",
                                        "description": "Additional data for tables or KPIs",
                                    },
                                },
                                "required": ["title"],
                            },
                        },
                        "filename": {
                            "type": "string",
                            "description": "Output filename (without extension)",
                        },
                    },
                    "required": ["title", "slides", "filename"],
                },
            },
            {
                "name": "create_kpi_slide",
                "description": "Create a standalone KPI summary slide as a presentation.",
                "input_schema": {
                    "type": "object",
                    "properties": {
                        "title": {"type": "string", "description": "Slide title"},
                        "kpis": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "name": {"type": "string"},
                                    "value": {"type": "string"},
                                    "trend": {
                                        "type": "string",
                                        "enum": ["up", "down", "flat"],
                                    },
                                    "change": {"type": "string"},
                                },
                                "required": ["name", "value"],
                            },
                            "description": "Array of KPI data objects",
                        },
                        "filename": {"type": "string"},
                    },
                    "required": ["title", "kpis", "filename"],
                },
            },
        ]

    def create_presentation(
        self,
        title: str,
        slides: list[dict],
        filename: str,
        subtitle: str = "",
    ) -> dict:
        """Create a full PowerPoint presentation."""
        if not PPTX_AVAILABLE:
            return self._create_presentation_json(title, slides, filename, subtitle)

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        # Title slide
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = title
        if slide.placeholders[1]:
            slide.placeholders[1].text = subtitle or datetime.now().strftime(
                "%B %d, %Y"
            )

        # Content slides
        for slide_data in slides:
            stype = slide_data.get("slide_type", "content")
            if stype == "table" and slide_data.get("data"):
                self._add_table_slide(prs, slide_data)
            elif stype == "kpi" and slide_data.get("data"):
                self._add_kpi_slide(prs, slide_data)
            elif stype == "two_column":
                self._add_two_column_slide(prs, slide_data)
            else:
                self._add_content_slide(prs, slide_data)

        filepath = self.output_dir / f"{filename}.pptx"
        prs.save(str(filepath))
        return {
            "status": "created",
            "filepath": str(filepath),
            "slide_count": len(prs.slides),
        }

    def create_kpi_slide(
        self, title: str, kpis: list[dict], filename: str
    ) -> dict:
        """Create a KPI-focused presentation."""
        slides = [
            {
                "title": title,
                "slide_type": "kpi",
                "data": {"kpis": kpis},
            }
        ]
        return self.create_presentation(
            title=title,
            slides=slides,
            filename=filename,
            subtitle="KPI Dashboard",
        )

    def _add_content_slide(self, prs: "Presentation", data: dict) -> None:
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = data.get("title", "")
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.text = data.get("content", "")

    def _add_two_column_slide(self, prs: "Presentation", data: dict) -> None:
        slide_layout = prs.slide_layouts[3]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = data.get("title", "")
        content = data.get("content", "")
        parts = content.split("|||")
        if len(parts) >= 2 and len(slide.placeholders) > 2:
            slide.placeholders[1].text = parts[0].strip()
            slide.placeholders[2].text = parts[1].strip()

    def _add_table_slide(self, prs: "Presentation", data: dict) -> None:
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = data.get("title", "")

        table_data = data.get("data", {})
        headers = table_data.get("headers", [])
        rows = table_data.get("rows", [])

        if not headers or not rows:
            return

        num_rows = len(rows) + 1
        num_cols = len(headers)
        left = Inches(0.5)
        top = Inches(2.0)
        width = Inches(12.0)
        height = Inches(0.4) * num_rows

        table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
        table = table_shape.table

        for i, header in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = str(header)

        for row_idx, row in enumerate(rows):
            for col_idx, val in enumerate(row):
                cell = table.cell(row_idx + 1, col_idx)
                cell.text = str(val)

    def _add_kpi_slide(self, prs: "Presentation", data: dict) -> None:
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = data.get("title", "")

        kpis = data.get("data", {}).get("kpis", [])
        num_kpis = len(kpis)
        if num_kpis == 0:
            return

        card_width = Inches(2.8)
        card_height = Inches(2.0)
        margin = Inches(0.3)
        start_left = Inches(0.5)
        top = Inches(2.5)

        for i, kpi in enumerate(kpis):
            left = start_left + i * (card_width + margin)
            shape = slide.shapes.add_shape(
                1, left, top, card_width, card_height  # Rectangle
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0x00, 0x3D, 0x6B)

            tf = shape.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = kpi.get("name", "")
            p.font.size = Pt(14)
            p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            p.alignment = PP_ALIGN.CENTER

            p2 = tf.add_paragraph()
            p2.text = kpi.get("value", "")
            p2.font.size = Pt(28)
            p2.font.bold = True
            p2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            p2.alignment = PP_ALIGN.CENTER

            trend = kpi.get("trend", "flat")
            change = kpi.get("change", "")
            arrow = {"up": "^", "down": "v", "flat": "-"}.get(trend, "-")
            p3 = tf.add_paragraph()
            p3.text = f"{arrow} {change}" if change else ""
            p3.font.size = Pt(12)
            color = (
                RGBColor(0x00, 0xFF, 0x00)
                if trend == "up"
                else RGBColor(0xFF, 0x00, 0x00)
                if trend == "down"
                else RGBColor(0xFF, 0xFF, 0xFF)
            )
            p3.font.color.rgb = color
            p3.alignment = PP_ALIGN.CENTER

    def _create_presentation_json(
        self, title: str, slides: list[dict], filename: str, subtitle: str
    ) -> dict:
        """Fallback: save presentation data as JSON when python-pptx is unavailable."""
        data = {
            "title": title,
            "subtitle": subtitle,
            "created_at": datetime.now().isoformat(),
            "slides": slides,
        }
        filepath = self.output_dir / f"{filename}.json"
        filepath.write_text(json.dumps(data, indent=2))
        return {
            "status": "created_json_fallback",
            "filepath": str(filepath),
            "slide_count": len(slides),
            "note": "python-pptx not installed; saved as JSON structure",
        }

    def handle_tool_call(self, tool_name: str, tool_input: dict) -> Any:
        dispatch = {
            "create_presentation": lambda: self.create_presentation(**tool_input),
            "create_kpi_slide": lambda: self.create_kpi_slide(**tool_input),
        }
        handler = dispatch.get(tool_name)
        if handler:
            return handler()
        return {"error": f"Unknown presentation tool: {tool_name}"}
